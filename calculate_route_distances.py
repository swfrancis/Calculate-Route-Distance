"""
Route Distance Calculator using OpenRouteService API.
Run this locally (not in a sandboxed environment) for accurate driving distances.

Prerequisites:
    pip install openpyxl requests

Usage:
    python calculate_route_distances.py

    Set your ORS API key via environment variable or edit ORS_API_KEY below:
        export ORS_API_KEY="your_key_here"
        python calculate_route_distances.py

Input files (must be in same directory):
    - multi-drops.xlsx          (original spreadsheet)
    - address_pairs.csv         (generated address pairs)

Output:
    - multi-drops-with-distances.xlsx  (updated spreadsheet with column J filled)
"""

import csv
import json
import math
import os
import re
import shutil
import time
from pathlib import Path

import requests
from openpyxl import load_workbook

# --- Config ---
ORS_API_KEY = os.environ.get(
    "ORS_API_KEY",
    "eyJvcmciOiI1YjNjZTM1OTc4NTExMTAwMDFjZjYyNDgiLCJpZCI6IjAwODFhZTRkYWNjYjRiZWFhNGUxOTJlNzhiYzljNGRjIiwiaCI6Im11cm11cjY0In0="
)
ORS_DIRECTIONS_URL = "https://api.openrouteservice.org/v2/directions/driving-car"
ORS_GEOCODE_URL = "https://api.openrouteservice.org/geocode/search"

GEOCODE_CACHE_FILE = "geocode_cache.json"
DISTANCE_CACHE_FILE = "distance_cache.json"
INPUT_XLSX = "multi-drops.xlsx"
PAIRS_CSV = "address_pairs.csv"
OUTPUT_XLSX = "multi-drops-with-distances.xlsx"

# ORS free tier: 40 directions/min, 2000/day
ORS_DIRECTIONS_DELAY = 1.6  # seconds between direction requests (~37/min)
ORS_GEOCODE_DELAY = 0.5     # seconds between geocode requests


# --- Cache helpers ---
def load_cache(path):
    if Path(path).exists():
        with open(path) as f:
            return json.load(f)
    return {}


def save_cache(data, path):
    with open(path, "w") as f:
        json.dump(data, f, indent=2)


# --- Geocoding via ORS Pelias ---
def geocode_address(address, cache):
    """Geocode an address using OpenRouteService's geocoding API."""
    cache_key = address.strip().upper()
    if cache_key in cache:
        return cache[cache_key]

    # Clean address for geocoding
    clean = address.strip()

    # Handle LOT prefix addresses
    if clean.upper().startswith("LOT "):
        m = re.match(r'LOT\s+\d+\s*\((\d+)\)\s*(.*)', clean, re.IGNORECASE)
        if m:
            clean = f"{m.group(1)} {m.group(2)}"
        else:
            clean = re.sub(r'^LOT\s+\d+\s*', '', clean, flags=re.IGNORECASE)

    # Handle unit/shop prefixes
    clean = re.sub(r'^(UNIT|SHOP|SUITE|LEVEL)\s+\d+[A-Z]?\s*[,/]?\s*', '', clean, flags=re.IGNORECASE)

    query = f"{clean}, Australia"

    headers = {"Authorization": ORS_API_KEY}
    params = {
        "api_key": ORS_API_KEY,
        "text": query,
        "boundary.country": "AU",
        "size": 1,
    }

    try:
        resp = requests.get(ORS_GEOCODE_URL, params=params, headers=headers, timeout=15)
        resp.raise_for_status()
        data = resp.json()
        features = data.get("features", [])
        if features:
            coords = features[0]["geometry"]["coordinates"]  # [lon, lat]
            result = (coords[1], coords[0])  # (lat, lon)
            cache[cache_key] = result
            return result

        # Fallback: try suburb + state + postcode only
        parts = address.split(",")
        if len(parts) >= 2:
            fallback_query = parts[-1].strip() + ", Australia"
            params["text"] = fallback_query
            resp = requests.get(ORS_GEOCODE_URL, params=params, headers=headers, timeout=15)
            resp.raise_for_status()
            data = resp.json()
            features = data.get("features", [])
            if features:
                coords = features[0]["geometry"]["coordinates"]
                result = (coords[1], coords[0])
                cache[cache_key] = result
                return result
    except Exception as e:
        print(f"  Geocode error for '{address}': {e}")

    print(f"  WARNING: Could not geocode '{address}'")
    cache[cache_key] = None
    return None


# --- Driving distance via ORS Directions ---
def get_driving_distance_km(coord1, coord2, dist_cache):
    """Get driving distance in km between two (lat, lon) coords using ORS Directions API."""
    cache_key = f"{coord1[0]:.6f},{coord1[1]:.6f}|{coord2[0]:.6f},{coord2[1]:.6f}"
    if cache_key in dist_cache:
        return dist_cache[cache_key]

    # ORS expects POST with coordinates as [[lon,lat],[lon,lat]]
    headers = {
        "Authorization": ORS_API_KEY,
        "Content-Type": "application/json",
    }
    body = {
        "coordinates": [
            [coord1[1], coord1[0]],  # [lon, lat]
            [coord2[1], coord2[0]],
        ],
    }

    try:
        resp = requests.post(ORS_DIRECTIONS_URL, json=body, headers=headers, timeout=30)
        resp.raise_for_status()
        data = resp.json()
        routes = data.get("routes", [])
        if routes:
            dist_m = routes[0]["summary"]["distance"]
            dist_km = round(dist_m / 1000.0, 1)
            dist_cache[cache_key] = dist_km
            return dist_km
        else:
            print(f"  ORS no route found")
    except requests.exceptions.HTTPError as e:
        if e.response is not None and e.response.status_code == 429:
            print("  Rate limited! Waiting 60s...")
            time.sleep(60)
            return get_driving_distance_km(coord1, coord2, dist_cache)
        print(f"  ORS error: {e}")
    except Exception as e:
        print(f"  ORS error: {e}")

    dist_cache[cache_key] = None
    return None


# --- Haversine fallback ---
def haversine_km(coord1, coord2):
    R = 6371.0
    lat1, lon1 = math.radians(coord1[0]), math.radians(coord1[1])
    lat2, lon2 = math.radians(coord2[0]), math.radians(coord2[1])
    dlat = lat2 - lat1
    dlon = lon2 - lon1
    a = math.sin(dlat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon / 2) ** 2
    return R * 2 * math.asin(math.sqrt(a))


# --- Main ---
def main():
    print("=" * 60)
    print("Route Distance Calculator (OpenRouteService)")
    print("=" * 60)

    if not ORS_API_KEY:
        print("ERROR: No API key. Set ORS_API_KEY environment variable.")
        return

    # Load address pairs
    pairs = []
    with open(PAIRS_CSV) as f:
        reader = csv.DictReader(f)
        for row in reader:
            pairs.append(row)

    load_numbers = set(r["load_number"] for r in pairs)
    print(f"\nLoaded {len(pairs)} address pairs across {len(load_numbers)} loads")

    # Collect unique addresses
    all_addresses = set()
    for row in pairs:
        all_addresses.add(row["from_address"])
        all_addresses.add(row["to_address"])
    print(f"Unique addresses to geocode: {len(all_addresses)}")

    # Step 1: Geocode all addresses
    print("\n--- Step 1: Geocoding addresses ---")
    geo_cache = load_cache(GEOCODE_CACHE_FILE)

    cached = sum(1 for a in all_addresses if a.strip().upper() in geo_cache)
    to_geocode = len(all_addresses) - cached
    print(f"Already cached: {cached}/{len(all_addresses)}")
    if to_geocode > 0:
        print(f"Need to geocode: {to_geocode} addresses")

    count = 0
    for addr in sorted(all_addresses):
        if addr.strip().upper() in geo_cache:
            continue
        count += 1
        print(f"  [{count}/{to_geocode}] Geocoding: {addr[:70]}...")
        geocode_address(addr, geo_cache)
        time.sleep(ORS_GEOCODE_DELAY)
        if count % 20 == 0:
            save_cache(geo_cache, GEOCODE_CACHE_FILE)

    save_cache(geo_cache, GEOCODE_CACHE_FILE)

    failed_geo = [a for a in all_addresses if geo_cache.get(a.strip().upper()) is None]
    if failed_geo:
        print(f"\nWARNING: {len(failed_geo)} addresses could not be geocoded:")
        for a in sorted(failed_geo):
            print(f"  - {a}")

    # Step 2: Get driving distances for each leg
    print("\n--- Step 2: Fetching driving distances via ORS ---")
    dist_cache = load_cache(DISTANCE_CACHE_FILE)

    cached_dists = 0
    for row in pairs:
        c1 = geo_cache.get(row["from_address"].strip().upper())
        c2 = geo_cache.get(row["to_address"].strip().upper())
        if c1 and c2:
            cache_key = f"{c1[0]:.6f},{c1[1]:.6f}|{c2[0]:.6f},{c2[1]:.6f}"
            if cache_key in dist_cache:
                cached_dists += 1

    to_fetch = len(pairs) - cached_dists
    print(f"Already cached: {cached_dists}/{len(pairs)}")
    if to_fetch > 0:
        est_minutes = (to_fetch * ORS_DIRECTIONS_DELAY) / 60
        print(f"Need to fetch: {to_fetch} distances (est. {est_minutes:.1f} min)")

    results = []
    fetch_count = 0
    for idx, row in enumerate(pairs):
        coord1 = geo_cache.get(row["from_address"].strip().upper())
        coord2 = geo_cache.get(row["to_address"].strip().upper())

        if coord1 is None or coord2 is None:
            print(f"  Load {row['load_number']} stop {row['from_stop']}->{row['to_stop']}: SKIP (geocode failed)")
            results.append(None)
            continue

        # Check if same location (< 50m apart)
        if haversine_km(coord1, coord2) < 0.05:
            results.append(0.0)
            continue

        cache_key = f"{coord1[0]:.6f},{coord1[1]:.6f}|{coord2[0]:.6f},{coord2[1]:.6f}"
        if cache_key in dist_cache:
            results.append(dist_cache[cache_key])
            continue

        fetch_count += 1
        dist = get_driving_distance_km(coord1, coord2, dist_cache)
        if dist is not None:
            print(f"  [{fetch_count}/{to_fetch}] Load {row['load_number']} "
                  f"stop {row['from_stop']}->{row['to_stop']}: {dist} km")
        else:
            print(f"  [{fetch_count}/{to_fetch}] Load {row['load_number']} "
                  f"stop {row['from_stop']}->{row['to_stop']}: FAILED")
        results.append(dist)
        time.sleep(ORS_DIRECTIONS_DELAY)

        if fetch_count % 30 == 0:
            save_cache(dist_cache, DISTANCE_CACHE_FILE)

    save_cache(dist_cache, DISTANCE_CACHE_FILE)

    # Build lookup: (load_number, from_stop, to_stop) -> distance
    leg_lookup = {}
    for row, dist in zip(pairs, results):
        key = (int(row["load_number"]), str(row["from_stop"]), str(row["to_stop"]))
        leg_lookup[key] = dist

    # Step 3: Write to spreadsheet
    print("\n--- Step 3: Writing results to spreadsheet ---")
    shutil.copy(INPUT_XLSX, OUTPUT_XLSX)
    wb = load_workbook(OUTPUT_XLSX)
    ws = wb.active

    current_load = None
    cumulative = 0.0
    prev_stop = None
    col_j = 10

    for row_idx in range(2, ws.max_row + 1):
        load_num = ws.cell(row=row_idx, column=1).value
        stop_order = ws.cell(row=row_idx, column=2).value

        if load_num is None:
            continue

        load_num_int = int(load_num)
        stop_str = str(stop_order)

        if stop_str == "DEPOT":
            current_load = load_num_int
            cumulative = 0.0
            prev_stop = "DEPOT"
            ws.cell(row=row_idx, column=col_j).value = 0
        else:
            key = (current_load, prev_stop, stop_str)
            leg_dist = leg_lookup.get(key)
            if leg_dist is not None:
                cumulative += leg_dist
                ws.cell(row=row_idx, column=col_j).value = round(cumulative, 1)
            else:
                ws.cell(row=row_idx, column=col_j).value = "N/A"
            prev_stop = stop_str

    wb.save(OUTPUT_XLSX)
    print(f"\nDone! Output saved to: {OUTPUT_XLSX}")

    success = sum(1 for r in results if r is not None)
    print(f"\nSummary: {success}/{len(results)} legs calculated successfully")
    if failed_geo:
        print(f"         {len(failed_geo)} addresses failed geocoding")


if __name__ == "__main__":
    main()
